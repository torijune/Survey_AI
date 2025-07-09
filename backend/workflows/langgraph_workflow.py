import asyncio
import json
import io
import re
from typing import Dict, Any, Optional, List, Union
import pandas as pd
import numpy as np
from langchain_openai import ChatOpenAI
from langchain.schema import HumanMessage, SystemMessage
import os
from dotenv import load_dotenv
import openpyxl
from scipy import stats
from collections import defaultdict
import openai

load_dotenv()

class AgentState:
    """Agent 상태를 관리하는 클래스"""
    def __init__(self, **kwargs):
        self.query = kwargs.get("query", "")
        self.file_path = kwargs.get("file_path", "")
        self.analysis_type = kwargs.get("analysis_type", True)
        self.selected_question = kwargs.get("selected_question", "")
        self.selected_table = kwargs.get("selected_table", None)
        self.selected_key = kwargs.get("selected_key", "")
        self.anchor = kwargs.get("anchor", [])
        self.linearized_table = kwargs.get("linearized_table", "")
        self.numeric_analysis = kwargs.get("numeric_analysis", "")
        self.table_analysis = kwargs.get("table_analysis", "")
        self.revised_analysis = kwargs.get("revised_analysis", "")
        self.hallucination_check = kwargs.get("hallucination_check", "")
        self.hallucination_reject_num = kwargs.get("hallucination_reject_num", 0)
        self.feedback = kwargs.get("feedback", "")
        self.polishing_result = kwargs.get("polishing_result", "")
        self.generated_hypotheses = kwargs.get("generated_hypotheses", "")
        self.uploaded_file = kwargs.get("uploaded_file", None)
        self.raw_data_file = kwargs.get("raw_data_file", None)
        self.raw_data = kwargs.get("raw_data", None)
        self.raw_variables = kwargs.get("raw_variables", None)
        self.raw_code_guide = kwargs.get("raw_code_guide", None)
        self.raw_question = kwargs.get("raw_question", None)
        self.test_type = kwargs.get("test_type", None)
        self.ft_test_result = kwargs.get("ft_test_result", {})
        self.ft_test_summary = kwargs.get("ft_test_summary", "")
        self.ft_error = kwargs.get("ft_error", None)
        self.lang = kwargs.get("lang", "한국어")
        self.tables = kwargs.get("tables", {})
        self.question_texts = kwargs.get("question_texts", {})
        self.question_keys = kwargs.get("question_keys", [])
        self.revised_analysis_history = kwargs.get("revised_analysis_history", [])
        self.user_id = kwargs.get("user_id", None)
        self.survey_data = kwargs.get("survey_data", None)

class LangGraphWorkflow:
    TABLE_ANALYSIS_PROMPT = {
        "한국어": """
당신은 통계 데이터를 바탕으로 인구집단 간 경향을 요약하는 데이터 분석 전문가입니다.
특히 아래 두 가지 정보를 종합적으로 활용해 핵심적 경향을 도출해야 합니다:

1️⃣ **통계 분석 결과(ft_test_summary)**: 어떤 대분류(row)가 통계적으로 유의미한 차이를 보였는지 알려줍니다.
2️⃣ **중요 변수(anchor)**: 어떤 column(항목)이 응답자의 전체 응답에서 가장 투표율이 높은 항목인지를 나타냅니다.

---

📝 설문 조사 질문:
{selected_question}

📊 표 데이터 (선형화된 형태):
{linearized_table}

📈 주요 항목 (변수들 중 가장 투표율이 높은 변수):
{anchor}

📈 통계 분석 결과 (통계적으로 유의미한 대분류):
{ft_test_summary}

---

⚠️ 참고: 만약 통계 분석 결과가 존재하지 않거나 사용자가 분석을 진행하지 않기로 선택한 경우, 주요 항목(anchor)을 중심으로 경향을 파악하고 이를 기반으로 요약할 것.

Let's think step by step.

🎯 분석 및 요약 지침:
1. 반드시 **F/T test 결과에서 통계적으로 유의미한 대분류만을 중심으로 분석**할 것 (p-value < 0.05, 유의성 별(*) 존재)
2. 모든 대분류 / 소분류를 나열하지 말고, **통계 분석 결과**에서 차이가 크고 의미 있는 대분류만 선택적으로 언급할 것
    - 통계적으로 유의미한 대분류가 없을 경우 (유의성 별(*)가 없을 경우) 주어진 p-value가 작은 대분류에서 주요 항목에 포함되는 대분류만 언급할 것
3. **절대 해석하지 말 것**. 수치적 차이에 대한 인과 해석(예: 건강에 민감해서, 주변에 있어서 등)은 모두 **금지**함
4. 외부 배경지식, 주관적 추론, 해석적 언급은 절대 금지. 표로부터 직접 **확인 가능한 사실만 서술**할 것
5. 수치 기반 경향을 다음과 같은 형식으로 서술하며 음슴체로 작성할 것 (예: ~했음, ~로 나타났음):
   - 상대적으로 더 높은 경향 보였음
   - 낮은 값을 나타냈음
6. 문장 간 연결어를 활용해 자연스럽게 서술하고, 너무 단조롭거나 반복적인 표현 (~했음. ~했음.)은 연속적으로 사용하지 말 것
7. **유의성이 없거나, 검정에서 제외된 항목은 절대 언급하지 말 것**
8. 모든 대분류가 유의성이 있고 중요하다면 모든 변수에 대해 설명하지 말고, **모든 대분류들이 중요했다고만 언급**할 것
9. **특정 대분류가 가장 두드러진 차이를 보였을 경우**, 해당 경향을 강조할 것
10. 숫자값을 직접 쓰지 말고 상대적인 경향만 언급할 것
""",
        "English": """
You are a data analyst summarizing trends across population groups based on statistical data.
You must integrate the following two pieces of information to identify key patterns:

1️⃣ **Statistical Test Results (ft_test_summary)**: indicates which row groups (categories) showed statistically significant differences
2️⃣ **Key Variables (anchor)**: columns (features) that had the highest overall selection rate among respondents

---

📝 Survey Question:
{selected_question}

📊 Table Data (Linearized):
{linearized_table}

📈 Key Variables (most frequently selected):
{anchor}

📈 Statistical Test Results (significant groups):
{ft_test_summary}

---

⚠️ Note: If there are no statistical results or if the user has opted out of statistical analysis, summarize based on key variables (anchor) and observed trends around them.

Let's think step by step.

🎯 Guidelines for Analysis and Summary:
1. Focus only on row groups that are statistically significant (p-value < 0.05, marked with asterisk)
2. Do not list all groups/subgroups; highlight only those with major, meaningful differences
    - If there are no statistically significant categories (if there are no significant stars (*)), only mention the categories included in the main items with small p-values.
3. **Do not interpret causality** (e.g., due to health sensitivity, etc.) – strictly prohibited
4. No external knowledge or speculation – write only what is verifiable from the table
5. Describe trends using expressions like:
   - Showed relatively higher trend
   - Showed lower values
6. Write in bullet-style declarative tone (e.g., “~was observed”, “~was shown”)
7. Use transition words to make the sentences flow naturally; avoid repetitive sentence endings
8. **Do not mention non-significant or excluded categories**
9. **If a particular group showed the strongest difference**, emphasize it
10. Do not mention actual numerical values, only describe relative trends
"""
    }
    HALLUCINATION_CHECK_PROMPT = {
        "한국어": """
당신은 통계 해석 결과를 검증하는 전문가입니다.

아래의 테이블 데이터와 수치 분석 결과(F/T-test 기반), 그리고 해당 결과를 바탕으로 작성된 요약 보고서가 주어집니다.

📝 설문 문항:
{selected_question}

📊 선형화된 테이블:
{linearized_table}

📈 수치 분석 결과 (F/T-test 결과 요약):
{ft_test_summary}

🧾 생성된 요약:
{table_analysis}

---

이 요약이 위의 수치 분석 결과를 **정확하고 일관성 있게** 반영하고 있는지 평가해주세요.

⚠️ 주의 사항 (위반 시 우선 피드백 제공, 심각한 왜곡에 한해 reject):
1. F/T-test에서 통계적으로 유의미한 차이가 확인된 대분류가 요약에 언급되지 않은 경우
2. 유의미한 차이가 확인된 대분류에서의 주요 경향이나 수치 결과가 왜곡되어 해석된 경우 (e.g. 더 높지 않은데 더 높다고 잘못 된 주장을 하는 경우)

🎯 평가 방식:
- 요약이 전체적으로 신뢰할 만하고 통계 결과를 잘 반영하면 "accept"
- 위 항목 위반 시 "reject: [이유]" 형식으로 출력

※ F/T-test 결과는 중요한 기준이지만, 사소한 누락은 reject 대신 피드백으로 처리해도 됩니다. 명백한 왜곡이나 중대한 누락 시에만 reject 하세요.
""",
        "English": """
You are a statistical analysis auditor.

Below is a statistical summary table (linearized format), F/T-test results, and a summary report written based on them.

📝 Survey question:
{selected_question}

📊 Linearized Table:
{linearized_table}

📈 Statistical Test Summary (F/T-test):
{ft_test_summary}

🧾 Generated Summary Report:
{table_analysis}

---

Please evaluate whether the summary accurately and consistently reflects the statistical test results above.

⚠️ Evaluation Guidelines (Provide feedback first. Only reject in cases of serious distortion):
1. If a major category with statistically significant difference is missing in the summary
2. If the key trends or directions are misinterpreted (e.g. stating it’s higher when it isn’t)

🎯 Evaluation Instructions:
- If the summary is overall reliable and reflects the results well, answer "accept"
- If any violations occur, return: "reject: [reason]"

※ F/T-test is a key basis, but minor omissions can be handled with feedback only. Use "reject" only for clear distortions or major omissions.
"""
    }
    REVISION_PROMPT = {
        "한국어": """
당신은 통계 데이터를 바탕으로 인구집단 간 패턴과 경향성을 객관적으로 요약하는 데이터 분석 전문가입니다.

아래는 테이블 분석 결과에 대해 일부 잘못된 해석이 포함된 요약입니다. 피드백과 사전에 생성된 가설을 참고하여 잘못된 내용을 제거하고, 원본 데이터를 기반으로 수치 기반의 객관적 분석을 다시 작성할 것.

📊 표 데이터 (선형화된 형태):
{linearized_table}

📈 주요 항목 (변수들 중 가장 투표율이 높은 변수):
{anchor}

📈 통계 분석 결과 (통계적으로 유의미한 대분류):
{ft_test_summary}

📝 Reject된 보고서 (수정해야할 보고서):
{report_to_modify}

❗ 피드백 (수정이 필요한 이유 또는 잘못된 부분):
{feedback}

---

Let's think step by step

🎯 수정 및 재작성 지침:

1. 반드시 **F/T test 결과에서 통계적으로 유의미한 대분류만을 중심으로 분석**할 것 (p-value < 0.05, 유의성 별(*) 존재)
2. 모든 대분류 / 소분류를 나열하지 말고, **검정 결과에서 차이가 크고 의미 있는 대분류만 선택적으로 언급**할 것
3. **절대 해석하지 말 것**. 수치적 차이에 대한 인과 해석(예: 건강에 민감해서, 주변에 있어서 등)은 모두 금지함
4. 외부 배경지식, 주관적 추론, 해석적 언급은 절대 금지. **표로부터 직접 확인 가능한 사실만 서술**할 것
5. 수치 기반 경향을 다음과 같은 형식으로 서술할 것:
   - 상대적으로 더 높은 경향 보였음
   - 낮은 값을 나타냈음
6. 보고서 음슴체로 작성할 것 (예: ~했음, ~로 나타났음)
7. 문장 간 연결어를 활용해 자연스럽게 서술하고, 너무 단조롭거나 반복적인 표현 (~했음. ~했음.)은 피할 것
8. **유의성이 없거나, 검정에서 제외된 항목은 절대 언급하지 말 것**
9. **특정 대분류가 가장 두드러진 차이를 보였을 경우**, 해당 경향을 강조할 것
10. 숫자값을 직접 쓰지 말고 상대적인 경향만 언급할 것
11. 이전 수정 버전의 문장 표현을 재사용하지 않고, 새로운 어휘와 구조로 작성할 것
12. 추론 과정을 작성하지 말고 최종적으로 수정한 보고서만 출력하세요.
""",
        "English": """
You are a data analyst who objectively summarizes population-level patterns based on statistical data.

Below is a summary that contains partially incorrect interpretations of a statistical table analysis. Based on the given feedback and hypotheses, revise the summary by removing inaccurate parts and rewrite a new objective analysis grounded in the data.

📊 Table data (linearized):
{linearized_table}

📈 Key variables (most selected):
{anchor}

📈 Statistical analysis results (significant categories):
{ft_test_summary}

📝 Rejected summary (needs revision):
{report_to_modify}

❗ Feedback (reason for revision or incorrect points):
{feedback}

---

Let's think step by step

🎯 Revision Guidelines:

1. Focus only on categories that showed statistically significant differences in the F/T test (p-value < 0.05, marked with *)
2. Do not list all categories/subcategories; mention only those with meaningful differences
3. **Do not provide causal interpretations** – explanations like “due to health concerns” or similar are prohibited
4. No external knowledge or speculation – write only what is verifiable from the table
5. Describe trends in a form such as:
   - Showed relatively higher trend
   - Showed lower values
6. Write in bullet-style declarative tone (e.g., “~was observed”, “~was shown”)
7. Use transition words to make the sentences flow naturally; avoid repetitive sentence endings
8. **Do not mention non-significant or excluded categories**
9. **If a particular group showed the strongest difference**, emphasize it
10. Do not explain the reasoning – only output the final revised summary
"""
    }
    POLISHING_PROMPT = {
        "한국어": """
당신은 통계 데이터를 바탕으로 작성된 한국어 보고서를 다듬는 문체 전문 에디터입니다.

아래는 통계 분석 결과를 요약한 초안입니다.  
문장이 단절적이거나(~했음. ~했음 반복), 표현이 중복되거나, 불필요한 인사이트가 포함되어 있다면, **의미를 변경하지 않고** 더 읽기 쉬운 문장으로 다듬으세요.

🎯 다음 지침을 엄격히 따르세요:

1. **내용 추가, 삭제 금지** — 수치 기반의 원문 정보에서 벗어나는 새로운 해석, 배경 설명, 인과관계 유추는 모두 금지
2. **'음슴체' 스타일 유지** — 예: ~했음, ~로 나타났음
3. **인사이트 제거** — ‘건강에 민감해서’, ‘직접 영향을 받아서’ 등 주관적 추론은 모두 제거하고, 표로부터 드러나는 사실만 유지
4. **통계적으로 유의한 항목(별표 포함된 대분류)**만 문장에 포함되었는지 확인할 것
5. **중복 표현 제거 및 연결** — 동일 의미의 표현 반복("높게 나타났음", "관심이 높았음")은 피하고 연결어를 통해 간결하게 정리
6. **단조로운 나열 피하기** — ~했음. ~했음. 반복하지 말고, 문장 구조를 다양화하고 연관된 항목은 한 문장으로 묶기
7. **다양한 표현 혼용** — 아래와 같은 표현을 적절히 섞어 사용할 것:
   - 두드러진 경향 보였음
   - 뚜렷한 차이를 나타냈음
   - 상대적으로 높은 값을 보였음
   - 가장 높게 확인됐음
8. **불필요한 소분류 또는 모든 그룹 나열 금지** — 요약은 특징적인 그룹 중심으로 간결하게 작성할 것
9. **표 기반 사실만 요약** — 수치 기반 경향만 전달하고, 해석은 포함하지 말 것
10. 통계 분석 결과, 성별과 연령대에서 통계적으로 유의미한 차이를 보였음. 이러한 문장처럼 통계 분석 결과를 직접 언급하지 말 것. 표에 나타난 수치 기반의 경향만 언급할 것.

아래의 기존 요약 중에서 위의 지침에서 어긋나는 문장이 있다면 지침을 따르도록 수정하세요. 하지만, 기존 요약에서 제공하는 값들은 그대로 유지할 것.

📝 기존 요약:
{raw_summary}

---

🎯 다듬어진 최종 요약문:
""",
        "English": """
You are a stylistic editor for statistical summaries written in Korean.

Below is a draft summary of a statistical analysis.  
If the sentences are too choppy ("~했음. ~했음." repetition), redundant, or include subjective insights, rewrite them into a more readable style **without altering their meaning**.

🎯 Strictly follow these instructions:

1. **No additions or deletions** — Do not add new interpretations, background, or causal reasoning beyond the original numeric-based content.
2. **Keep declarative tone** — Use styles like: "~was observed", "~was shown"
3. **Remove speculative insights** — Phrases like “due to health concerns” or “because they were affected” must be removed; stick only to observable facts
4. **Only include categories with statistical significance (asterisked)** in the report
5. **Eliminate and connect duplicates** — Avoid repeating the same idea (e.g., “was high”, “interest was high”); connect with transitions
6. **Avoid monotonous structure** — Don’t repeat "~was observed." repeatedly; vary structure and combine related findings into single sentences
7. **Use varied expressions** — Mix in phrases like:
   - Showed notable trend
   - Displayed clear difference
   - Exhibited relatively high values
   - Recorded the highest
8. **Avoid listing all subgroups** — Focus on concise summaries of characteristic groups
9. **Only report table-based facts** — Do not include interpretations; describe numeric trends only

📝 Original draft:
{raw_summary}

---

🎯 Polished final summary:
"""
    }
    def __init__(self):
        self.llm = ChatOpenAI(
            model="gpt-4o-mini",
            temperature=0.3,
            api_key=os.getenv("OPENAI_API_KEY")
        )
    
    def normalize_key(self, key: str) -> str:
        """키 정규화"""
        return key.strip().replace(' ', '').replace('-', '').replace('_', '').replace('.', '').upper()
    
    def find_matching_key(self, target_key: str, available_keys: List[str]) -> Optional[str]:
        """키 매칭 찾기"""
        normalized_target = self.normalize_key(target_key)
        
        # 정확한 매칭
        for key in available_keys:
            if self.normalize_key(key) == normalized_target:
                return key
        
        # 포함 관계 매칭
        for key in available_keys:
            normalized_key = self.normalize_key(key)
            if normalized_key in normalized_target or normalized_target in normalized_key:
                return key
        
        # 유사도 기반 매칭 (Levenshtein distance)
        best_match = None
        best_score = float('inf')
        
        for key in available_keys:
            normalized_key = self.normalize_key(key)
            distance = self.levenshtein_distance(normalized_target, normalized_key)
            max_length = max(len(normalized_target), len(normalized_key))
            similarity = distance / max_length
            
            if similarity < 0.3 and similarity < best_score:
                best_match = key
                best_score = similarity
        
        return best_match
    
    def levenshtein_distance(self, str1: str, str2: str) -> int:
        """Levenshtein 거리 계산"""
        matrix = [[0] * (len(str2) + 1) for _ in range(len(str1) + 1)]
        
        for i in range(len(str1) + 1):
            matrix[i][0] = i
        
        for j in range(len(str2) + 1):
            matrix[0][j] = j
        
        for i in range(1, len(str1) + 1):
            for j in range(1, len(str2) + 1):
                if str1[i-1] == str2[j-1]:
                    matrix[i][j] = matrix[i-1][j-1]
                else:
                    matrix[i][j] = min(
                        matrix[i-1][j-1] + 1,
                        matrix[i][j-1] + 1,
                        matrix[i-1][j] + 1
                    )
        
        return matrix[len(str1)][len(str2)]
    
    async def load_survey_tables(self, file_content: bytes, file_name: str, sheet_name: str = "통계표"):
        """설문 테이블 로드 - DataProcessor의 load_survey_tables와 동일하게 동작하도록 수정"""
        from utils.data_processor import DataProcessor
        data_processor = DataProcessor()
        return data_processor.load_survey_tables(file_content, file_name, sheet_name)
    
    def linearize_row_wise(self, table: pd.DataFrame) -> str:
        """테이블을 행 단위로 선형화"""
        if table is None or table.empty:
            return ""
        
        # 숫자 컬럼만 선택
        numeric_cols = table.select_dtypes(include=[np.number]).columns.tolist()
        if not numeric_cols:
            return ""
        
        # 첫 번째 컬럼 (그룹명)과 숫자 컬럼들만 사용
        result_cols = [table.columns[0]] + numeric_cols
        result_df = table[result_cols]
        
        # 행 단위로 문자열 변환
        rows = []
        for _, row in result_df.iterrows():
            row_str = " | ".join([str(val) for val in row.values])
            rows.append(row_str)
        
        return "\n".join(rows)
    
    async def make_openai_call(self, messages: List[Dict[str, str]], model: str = "gpt-4o-mini", temperature: float = 0.3) -> str:
        """OpenAI API 호출"""
        try:
            langchain_messages = []
            for msg in messages:
                if msg["role"] == "system":
                    langchain_messages.append(SystemMessage(content=msg["content"]))
                else:
                    langchain_messages.append(HumanMessage(content=msg["content"]))
            
            response = await self.llm.ainvoke(langchain_messages)
            return response.content.strip()
        except Exception as e:
            raise Exception(f"OpenAI API 호출 실패: {str(e)}")
    
    async def table_parser_node(self, state: AgentState, on_step=None) -> AgentState:
        """테이블 파서 노드"""
        if on_step:
            on_step("📊 테이블 파서 노드 시작")
        
        # 파일에서 테이블 파싱
        if state.uploaded_file:
            parsed_data = await self.load_survey_tables(
                state.uploaded_file,
                state.file_path
            )
            state.tables = parsed_data["tables"]
            state.question_texts = parsed_data["question_texts"]
            state.question_keys = parsed_data["question_keys"]
        
        # 선택된 키 매칭
        if state.selected_key and state.tables:
            matching_key = self.find_matching_key(state.selected_key, list(state.tables.keys()))
            if matching_key:
                state.selected_key = matching_key
                state.selected_table = state.tables[matching_key]
                state.selected_question = state.question_texts[matching_key]
            elif state.analysis_type:
                raise Exception(f"선택된 질문 키 '{state.selected_key}'에 해당하는 테이블이 없습니다.")
            else:
                # 배치 분석의 경우 첫 번째 질문 사용
                state.selected_key = state.question_keys[0]
                state.selected_table = state.tables[state.selected_key]
                state.selected_question = state.question_texts[state.selected_key]
        
        # 테이블 선형화
        if state.selected_table is not None:
            state.linearized_table = self.linearize_row_wise(state.selected_table)
        
        return state
    
    async def hypothesis_generate_node(self, state: AgentState, on_step=None) -> AgentState:
        """가설 생성 노드 (Streamlit 로직 반영)"""
        if on_step:
            on_step("💡 가설 생성 노드 시작")

        selected_table = state.selected_table
        selected_question = state.selected_question
        lang = state.lang if hasattr(state, 'lang') else "한국어"

        # row와 column name 추출 (Streamlit 방식)
        import pandas as pd
        if selected_table is not None and isinstance(selected_table, pd.DataFrame):
            if "대분류" in selected_table.columns and "소분류" in selected_table.columns:
                selected_table = selected_table.copy()
                selected_table["row_name"] = selected_table["대분류"].astype(str) + "_" + selected_table["소분류"].astype(str)
                row_names = selected_table["row_name"].dropna().tolist()
            else:
                row_names = list(selected_table.index)
            column_names = list(selected_table.columns)
        else:
            row_names = []
            column_names = []

        row_names_str = ", ".join(map(str, row_names))
        column_names_str = ", ".join(map(str, column_names))

        # 프롬프트 정의 (Streamlit과 동일)
        HYPOTHESIS_PROMPT = {
            "한국어": """
당신은 통계 데이터를 해석하는 데이터 과학자입니다.
아래는 분석할 표의 row명 (index)과 column명입니다.

row: {row_names}
column: {column_names}

당신의 임무는, 사용자의 질문 ("{selected_question}")과 관련해  
데이터에서 확인할 수 있을 법한 가설(패턴, 관계)을 2~5개 정도 제안하는 것입니다.

예시:
1. 연령대가 높을수록 관심도가 높을 것이다.
2. 기저질환이 있는 경우 관심도가 높을 것이다.

- 데이터 기반으로 합리적인 가설만 생성할 것
- 외부 지식은 절대 사용 금지
- 문장 길이는 짧고, 번호 리스트로 작성
""",
            "English": """
You are a data scientist interpreting statistical tables.
Below are the row names (index) and column names of the table to be analyzed.

row: {row_names}
column: {column_names}

Your task is to propose 2 to 5 hypotheses (patterns or relationships) relevant to the user's question ("{selected_question}") that could be inferred from the data.

Example:
1. Older age groups may show higher interest.
2. Those with chronic illnesses may show higher concern.

- Only propose reasonable hypotheses based on the data
- Do not use external knowledge
- Keep sentences short and format as a numbered list
"""
        }

        prompt = HYPOTHESIS_PROMPT[lang].format(
            row_names=row_names_str,
            column_names=column_names_str,
            selected_question=selected_question
        )

        messages = [
            {"role": "system", "content": "당신은 통계 분석 전문가입니다." if lang == "한국어" else "You are a statistical analysis expert."},
            {"role": "user", "content": prompt}
        ]

        state.generated_hypotheses = await self.make_openai_call(messages)
        return state
    
    def rule_based_test_type_decision(self, question_text=""):
        """
        질문 텍스트에 복수응답/순위/다중 등 키워드가 있으면 manual, 아니면 None
        """
        multi_response_keywords = [
            "1+2", "1+2+3", "복수", "다중", "multiple", "rank", "ranking", "우선순위", "복수응답", "순위"
        ]
        text_to_check = question_text.lower()
        if any(keyword.lower() in text_to_check for keyword in multi_response_keywords):
            return "manual"
        return None

    async def decide_batch_test_types(self, question_infos, lang="한국어"):
        """
        전체 분석용: 여러 질문에 대해 통계 검정 방법을 일괄 결정하는 함수
        - 각 질문별로 복수응답/임의(manual) 여부는 rule-based로 판정
        - 복수응답이 아닌 질문들은 LLM에 한 번에 프롬프트로 전달하여 test_type(ft_test/chi_square) 결정
        - LLM 응답을 파싱해 question_key -> test_type 매핑 반환
        Args:
            question_infos: List[dict] (각 dict: {key, text, columns})
            lang: 언어 (default: 한국어)
        Returns:
            Dict[str, str] (question_key -> test_type)
        """
        # 1. manual 판정 (복수응답/순위/다중 등 키워드)
        manual_keys = set()
        non_manual_questions = []
        for q in question_infos:
            key, text, columns = q['key'], q['text'], q['columns']
            if self.rule_based_test_type_decision(text) == 'manual':
                manual_keys.add(key)
            else:
                non_manual_questions.append(q)
        # 2. LLM 프롬프트 생성 (복수응답 아닌 질문들)
        prompt_lines = []
        for q in non_manual_questions:
            col_str = ', '.join(q['columns'])
            prompt_lines.append(f"{q['key']}: {col_str}")
        prompt_body = '\n'.join(prompt_lines)
        if lang == "한국어":
            prompt = f"""
            아래는 설문 통계표의 각 질문별 열 이름 목록입니다.
            
            당신의 임무는 각 질문에 대해 **가장 적합한 통계 검정 방법**(ft_test 또는 chi_square)을 결정하는 것입니다.
            
            - ft_test: 평균, 점수, 비율 등 연속형(수치형) 데이터에 적합
            - chi_square: 항목 선택, 다중응답 등 범주형(선택형) 데이터에 적합
            
            아래 기준을 참고하세요:
            - 열 이름에 '평균', '점수', '%', '비율' 등이 포함되어 있으면 ft_test
            - 항목 선택, 다중응답, 범주형 선택지면 chi_square
            
            아래와 같은 형식으로만 답변하세요(설명 없이):
            
            예시:
            Q1: ft_test
            Q2: chi_square
            Q3: ft_test
            
            질문별 열 목록:
            {prompt_body}
            
            답변 형식:
            Q1: ft_test
            Q2: chi_square
            ...
            """
        else:
            prompt = f"""
            Below are the column headers for each survey question.
        
            Your task is to determine the **most appropriate statistical test type** (ft_test or chi_square) for each question.
            
            - ft_test: Use for continuous/numeric data (mean, score, %, ratio, etc.)
            - chi_square: Use for categorical/choice/multiple response data
            
            Guidelines:
            - If column names include 'mean', 'score', '%', 'ratio', etc. → ft_test
            - If the question is about selecting items, multiple responses, or categorical choices → chi_square
            - If the question is multiple response/ranking, use 'manual' (already auto-classified)
            
            Please answer in the following format (no explanation):
            
            Example:
            Q1: ft_test
            Q2: chi_square
            Q3: ft_test
            
            Question columns:
            {prompt_body}
            
            Answer format:
            Q1: ft_test
            Q2: chi_square
            ...
            """
        # 3. LLM 호출 (없으면 모두 ft_test)
        llm_result = None
        if non_manual_questions:
            messages = [
                {"role": "system", "content": "당신은 통계 전문가입니다." if lang == "한국어" else "You are a statistics expert."},
                {"role": "user", "content": prompt}
            ]
            try:
                llm_result = await self.make_openai_call(messages)
            except Exception as e:
                print(f"[decide_batch_test_types] LLM 호출 실패: {e}")
                # fallback: 모두 ft_test
                llm_result = '\n'.join([f"{q['key']}: ft_test" for q in non_manual_questions])
        # 4. LLM 응답 파싱
        test_type_map = {key: 'manual' for key in manual_keys}
        if llm_result:
            import re
            for line in llm_result.splitlines():
                m = re.match(r"([\w\-]+):\s*(ft_test|chi_square)", line.strip(), re.I)
                if m:
                    key, ttype = m.group(1), m.group(2).lower()
                    test_type_map[key] = ttype
        # 5. 누락된 질문은 기본값(ft_test)
        for q in question_infos:
            if q['key'] not in test_type_map:
                test_type_map[q['key']] = 'ft_test'
        return test_type_map

    async def test_decision_node(self, state: AgentState, on_step=None) -> AgentState:
        """통계 검정 방법 결정 노드 (질문 텍스트 manual 체크, 컬럼명 LLM 판정)"""
        if on_step:
            on_step("🧭 통계 검정 결정 노드 시작")
        if state.selected_table is not None:
            columns = state.selected_table.columns.tolist()
            IGNORE_COLUMNS = {"대분류", "소분류", "사례수", "row_name"}
            filtered_columns = [col for col in columns if col not in IGNORE_COLUMNS]
            question_text = getattr(state, 'selected_question', "")

            # manual 체크 (질문 텍스트만)
            manual_check = self.rule_based_test_type_decision(question_text)
            if manual_check == "manual":
                state.test_type = "manual"
                return state

            # LLM 프롬프트 (컬럼명만)
            TEST_TYPE_PROMPT = """
당신은 통계 전문가입니다.

아래는 설문 응답 결과 테이블의 열 이름 목록입니다. 이 열들은 응답자들이 선택하거나 평가한 설문 문항의 결과로 구성된 통계표입니다.

당신의 임무는, 이 테이블이 **어떤 통계 검정(F/T-test 또는 Chi-square)** 에 적합한지를 판단하는 것입니다.

📋 열 이름 목록:
{column_names}

---
Let's think step by step

판단 기준:

- `ft_test` (연속형 수치 응답):
    - 문항이 1~5점 척도, 평균, 비율, 점수 등 숫자 기반으로 요약되어 있다면 F-test 또는 T-test가 적절합니다.
    - 예시 열 이름: "평균", "만족도 점수", "~% 비율", "5점 척도", "평균 점수", "관심도 평균"
    - "전혀 관심이 없다", "매우 관심 있다" 등은 실제로는 선택지이지만, 빈도나 비율로 수치화되었을 경우 → 연속형으로 판단

- `chi_square` (범주형 선택 응답):
    - 문항이 응답자들이 특정 항목을 **선택**하거나 **다중선택**한 결과일 경우, 범주형 응답으로 보고 카이제곱 검정이 적합합니다.
    - 예시 열 이름: "주요 이용시설", "선택 이유", "가장 많이 선택한 장소", "다중 응답"
❗ 오판 주의:
- 응답 선택지 이름(예: "전혀 관심 없다", "매우 관심 있다")가 열 이름에 포함되더라도, **비율, 평균 등의 수치형 요약**이면 `ft_test`로 간주합니다.
- 테이블이 전체적으로 평균값 또는 %비율 중심이면 `ft_test` 선택이 더 적절합니다.

---

📌 답변 형식: 아래의 형식처럼 선택의 이유에 대해서 답변하지 말고 "적합한 통계 검정의 방법만" 출력하세요.

- 반드시 다음 중 하나로만 답해주세요 (소문자):
    - ft_test
    - chi_square

적합한 통계 방법: (ft_test 또는 chi_square)
"""
            column_names_str = ", ".join(filtered_columns)
            prompt = TEST_TYPE_PROMPT.format(column_names=column_names_str)

            # LLM 호출
            messages = [
                {"role": "system", "content": "당신은 통계 전문가입니다."},
                {"role": "user", "content": prompt}
            ]
            llm_result = await self.make_openai_call(messages)
            llm_output = llm_result.strip() if hasattr(llm_result, 'strip') else str(llm_result)

            # normalize
            def normalize_test_type(llm_output: str) -> str:
                if "chi" in llm_output.lower():
                    return "chi_square"
                elif "ft" in llm_output.lower():
                    return "ft_test"
                else:
                    return "unknown"
            state.test_type = normalize_test_type(llm_output)
        return state
    
    def assign_significance_stars(self, p_value):
        if p_value < 0.001:
            return "***"
        elif p_value < 0.01:
            return "**"
        elif p_value < 0.05:
            return "*"
        else:
            return ""

    def extract_demo_mapping_from_dataframe(self, df, column="Unnamed: 0"):
        print(f"[extract_demo_mapping_from_dataframe] called with column: {column}")
        col = df[column].dropna().astype(str).reset_index(drop=True)
        print(f"[extract_demo_mapping_from_dataframe] col values: {col.tolist()}")
        cut_idx = None
        for i, val in enumerate(col):
            if val.strip() == 'DEMO1':
                cut_idx = i
                break
        sliced = col[:cut_idx] if cut_idx is not None else col

        demo_dict = {}
        import re
        for entry in sliced:
            entry = str(entry).strip()
            match = re.match(r"(DEMO\d+)[\s'\"]+(.+?)[\'\"\s\.]*$", entry)
            if match:
                key = match.group(1)
                label = match.group(2).strip()
                demo_dict[key] = label
        print(f"[extract_demo_mapping_from_dataframe] demo_dict: {demo_dict}")
        return demo_dict

    def summarize_ft_test(self, result_df, lang: str = "한국어") -> str:
        if not isinstance(result_df, pd.DataFrame) or result_df.empty:
            return "통계 분석 결과가 없습니다."
        significant = result_df[result_df["유의성"] != ""]
        summary = []
        if not significant.empty:
            sig_items = significant["대분류"].tolist()
            if len(sig_items) == len(result_df):
                summary.append("모든 항목에서 통계적으로 유의미한 차이가 관찰되었음. 대분류 전반에 걸쳐 의미 있는 차이가 존재함." if lang == "한국어" else "All categories showed statistically significant differences. Broad variation was observed across major groups.")
            else:
                summary.append(f"{', '.join(sig_items)}는 통계적으로 유의한 차이를 보였음." if lang == "한국어" else f"{', '.join(sig_items)} showed statistically significant differences.")
        else:
            if not result_df.empty:
                top3 = result_df.nsmallest(3, "p-value")[["대분류", "p-value"]]
                top3_text = ", ".join(f"{row['대분류']} (p={row['p-value']})" for _, row in top3.iterrows())
                summary.append(f"통계적으로 유의한 항목은 없었지만, 상대적으로 p-value가 낮은 항목은 {top3_text} 순이었음." if lang == "한국어" else f"No items reached statistical significance, but the ones with the lowest p-values were: {top3_text}.")
        return "  ".join(summary)

    def run_statistical_tests(self, test_type, df, question_key, demo_dict):
        import pandas as pd
        import numpy as np
        import scipy.stats as stats
        print(f"[run_statistical_tests] test_type: {test_type}")
        print(f"[run_statistical_tests] question_key: {question_key}")
        print(f"[run_statistical_tests] demo_dict: {demo_dict}")
        print(f"[run_statistical_tests] df.columns: {df.columns.tolist()}")
        # FT-test
        def run_ft_test_df(df: pd.DataFrame, question_key: str, demo_dict: dict) -> pd.DataFrame:
            question_key = question_key.replace("-", "_").strip()
            rows = []
            for demo_col, label in demo_dict.items():
                print(f"  [FT] demo_col: {demo_col}, label: {label}")
                if demo_col not in df.columns:
                    print(f"    [FT] demo_col '{demo_col}' not in df.columns")
                    continue
                try:
                    groups = df.groupby(demo_col)[question_key].apply(list)
                    group_values = [pd.Series(values).dropna().tolist() for values in groups]
                    print(f"    [FT] group_values lens: {[len(g) for g in group_values]}")
                    if len(group_values) < 2:
                        print(f"    [FT] group_values < 2, skip")
                        continue
                    levene_stat, levene_p = stats.levene(*group_values)
                    if len(group_values) == 2:
                        test_stat, test_p = stats.ttest_ind(
                            group_values[0], group_values[1],
                            equal_var=(levene_p > 0.05)
                        )
                    else:
                        test_stat, test_p = stats.f_oneway(*group_values)
                    row = {
                        "대분류": label,
                        "통계량": round(abs(test_stat), 3),
                        "p-value": round(test_p, 4),
                        "유의성": self.assign_significance_stars(test_p)
                    }
                    print(f"    [FT] result row: {row}")
                    rows.append(row)
                except Exception as e:
                    print(f"    [FT] Exception: {e}")
                    continue
            result_df = pd.DataFrame(rows)
            print(f"[run_ft_test_df] result_df shape: {result_df.shape}")
            print(f"[run_ft_test_df] result_df: {result_df}")
            return result_df
        # Chi-square
        def run_chi_square_test_df(df: pd.DataFrame, question_key: str, demo_dict: dict) -> pd.DataFrame:
            question_key = question_key.replace("-", "_").strip()
            rows = []
            for demo_col, label in demo_dict.items():
                print(f"  [CHI] demo_col: {demo_col}, label: {label}")
                if demo_col not in df.columns:
                    print(f"    [CHI] demo_col '{demo_col}' not in df.columns")
                    continue
                try:
                    normalized_columns = {col.replace("-", "_").strip(): col for col in df.columns}
                    contingency_table = pd.crosstab(df[demo_col], df[normalized_columns[question_key]])
                    print(f"    [CHI] contingency_table shape: {contingency_table.shape}")
                    if contingency_table.shape[0] < 2 or contingency_table.shape[1] < 2:
                        print(f"    [CHI] contingency_table too small, skip")
                        continue
                    chi2, p, dof, expected = stats.chi2_contingency(contingency_table)
                    row = {
                        "대분류": label,
                        "통계량": round(chi2, 3),
                        "p-value": round(p, 4),
                        "유의성": self.assign_significance_stars(p)
                    }
                    print(f"    [CHI] result row: {row}")
                    rows.append(row)
                except Exception as e:
                    print(f"    [CHI] Exception: {e}")
                    continue
            result_df = pd.DataFrame(rows)
            print(f"[run_chi_square_test_df] result_df shape: {result_df.shape}")
            print(f"[run_chi_square_test_df] result_df: {result_df}")
            return result_df
        # manual
        def run_manual_analysis(df: pd.DataFrame, question_key: str, demo_dict: dict) -> pd.DataFrame:
            question_key = question_key.replace("-", "_").strip()
            try:
                overall_row = df[df["대분류"].astype(str).str.strip() == "전 체"]
                if overall_row.empty:
                    print("    [MANUAL] '전 체' 대분류 행이 존재하지 않습니다.")
                    return pd.DataFrame([])
                overall_value = overall_row[question_key].values[0]
                overall_n = overall_row["사례수"].values[0]
                overall_std = df[question_key].std()
                std_error = overall_std / np.sqrt(overall_n)
                z_score = 1.96
                ci_lower = overall_value - z_score * std_error
                ci_upper = overall_value + z_score * std_error
                rows = []
                for idx, row in df.iterrows():
                    if row["대분류"] == "전 체":
                        continue
                    group_value = row[question_key]
                    group_label = f"{row['대분류']} - {row['소분류']}" if pd.notna(row['소분류']) else row['대분류']
                    significant = group_value < ci_lower or group_value > ci_upper
                    rows.append({
                        "대분류": group_label,
                        "평균값": group_value,
                        "유의미 여부": "유의미함" if significant else "무의미함",
                        "기준 평균": overall_value,
                        "신뢰구간": f"{round(ci_lower,1)} ~ {round(ci_upper,1)}",
                        "유의성": "*" if significant else ""
                    })
                result_df = pd.DataFrame(rows)
                print(f"[run_manual_analysis] result_df shape: {result_df.shape}")
                print(f"[run_manual_analysis] result_df: {result_df}")
                return result_df
            except Exception as e:
                print(f"    [MANUAL] Exception: {e}")
                return pd.DataFrame([])
        if test_type == "ft_test":
            return run_ft_test_df(df, question_key, demo_dict)
        elif test_type =="chi_square":
            return run_chi_square_test_df(df, question_key, demo_dict)
        elif test_type == "manual":
            return run_manual_analysis(df, question_key, demo_dict)
        else:
            print(f"[run_statistical_tests] Invalid test_type: {test_type}")
            raise ValueError(f"❌ 잘못된 test_type: {test_type}")

    async def ft_analysis_node(self, state: AgentState, on_step=None) -> AgentState:
        if on_step:
            on_step("✅ F/T 분석 노드 시작")
        try:
            import pandas as pd
            from io import BytesIO
            if hasattr(state, 'raw_data_file') and state.raw_data_file is not None:
                print("[ft_analysis_node] raw_data_file exists")
                raw_data_file = BytesIO(state.raw_data_file) if isinstance(state.raw_data_file, (bytes, bytearray)) else state.raw_data_file
                raw_data = pd.read_excel(raw_data_file, sheet_name="DATA")
                demo_df = pd.read_excel(raw_data_file, sheet_name="DEMO")
                print(f"[ft_analysis_node] raw_data.columns: {raw_data.columns.tolist()}")
                print(f"[ft_analysis_node] demo_df.columns: {demo_df.columns.tolist()}")
                raw_data.columns = [col.replace("-", "_").strip() for col in raw_data.columns]
                demo_mapping = self.extract_demo_mapping_from_dataframe(demo_df)
                print(f"[ft_analysis_node] demo_mapping: {demo_mapping}")
                test_type = getattr(state, 'test_type', None)
                question_key = getattr(state, 'selected_key', None)
                lang = getattr(state, 'lang', "한국어")
                print(f"[ft_analysis_node] test_type: {test_type}, question_key: {question_key}, lang: {lang}")
                result_df = self.run_statistical_tests(test_type, raw_data, question_key, demo_mapping)
                print(f"[ft_analysis_node] result_df shape: {result_df.shape}")
                print(f"[ft_analysis_node] result_df: {result_df}")
                summary_text = self.summarize_ft_test(result_df, lang=lang)
                state.raw_data = raw_data
                state.ft_test_result = result_df
                state.ft_test_summary = summary_text
            else:
                print("[ft_analysis_node] raw_data_file is None")
                state.ft_error = "raw_data_file이 없습니다."
        except Exception as e:
            print(f"[ft_analysis_node] Exception: {e}")
            state.ft_error = str(e)
        return state
    
    async def get_anchor_node(self, state: AgentState, on_step=None) -> AgentState:
        """앵커 추출 노드 (Streamlit get_anchor 로직 반영)"""
        if on_step:
            on_step("📌 앵커 추출 노드 시작")
        import pandas as pd
        selected_table = getattr(state, 'selected_table', None)
        lang = getattr(state, 'lang', '한국어')
        anchor = []
        if selected_table is not None and isinstance(selected_table, pd.DataFrame):
            exclude_cols = set([
                "대분류", "소분류", "사례수",
                "대분류_소분류"
            ])
            def is_anchor_candidate(col):
                col_str = str(col).strip()
                if col_str in exclude_cols:
                    return False
                # %로 끝나는 컬럼 제외 (공백 포함)
                if col_str.rstrip().endswith('%'):
                    return False
                # 평균, 합계, std, score 등 포함시 제외
                if any(keyword in col_str for keyword in ['평균', '합계', 'std', 'score']):
                    return False
                return True
            # "전 체" 또는 "전체" row 찾기 (공백 제거)
            total_row = selected_table[selected_table["대분류"].astype(str).str.strip().str.replace(" ", "") == "전체"]
            if total_row.empty:
                state.anchor = []
                state.ft_error = "❌ '전 체' 또는 '전체'에 해당하는 행이 존재하지 않습니다."
                return state
            total_row = total_row.iloc[0]
            candidate_cols = [col for col in selected_table.columns if is_anchor_candidate(col)]
            values = total_row[candidate_cols]
            values_numeric = pd.to_numeric(values, errors='coerce')
            high_value_cols = values_numeric.dropna().sort_values(ascending=False)
            cumulative_sum = 0
            selected_cols = []
            for col, val in high_value_cols.items():
                cumulative_sum += val
                selected_cols.append(col)
                if cumulative_sum >= 60:
                    break
            anchor = selected_cols
        state.anchor = anchor
        return state
    
    async def table_analyzer(self, state: AgentState, on_step=None) -> AgentState:
        """테이블 분석 노드 (Streamlit 프롬프트/로직 반영)"""
        if on_step:
            on_step("🤖 테이블 분석 노드 시작")
        lang = getattr(state, 'lang', '한국어')
        linearized_table = getattr(state, 'linearized_table', '')
        ft_test_summary = getattr(state, 'ft_test_summary', '')
        selected_question = getattr(state, 'selected_question', '')
        anchor = getattr(state, 'anchor', '없음')
        prompt = self.TABLE_ANALYSIS_PROMPT[lang].format(
            selected_question=selected_question,
            linearized_table=linearized_table,
            ft_test_summary=str(ft_test_summary),
            anchor=anchor
        )
        messages = [
            {"role": "system", "content": "당신은 통계 분석 전문가입니다." if lang == "한국어" else "You are a statistical analysis expert."},
            {"role": "user", "content": prompt}
        ]
        state.table_analysis = await self.make_openai_call(messages)
        return state
    
    async def hallucination_check_node(self, state: AgentState, on_step=None) -> AgentState:
        """환각 검증 노드 (Streamlit 프롬프트/로직 반영)"""
        if on_step:
            on_step("🧠 환각 평가 노드 시작")
        lang = getattr(state, 'lang', '한국어')
        hallucination_reject_num = getattr(state, 'hallucination_reject_num', 0)
        # 수정 이력이 있으면 마지막 수정본 사용
        if hasattr(state, 'revised_analysis_history') and state.revised_analysis_history:
            table_analysis = state.revised_analysis_history[-1]
        else:
            table_analysis = getattr(state, 'table_analysis', '')
        prompt = self.HALLUCINATION_CHECK_PROMPT[lang].format(
            selected_question=getattr(state, 'selected_question', ''),
            linearized_table=getattr(state, 'linearized_table', ''),
            ft_test_summary=str(getattr(state, 'ft_test_summary', '')),
            table_analysis=table_analysis
        )
        messages = [
            {"role": "system", "content": "당신은 통계 해석 결과를 검증하는 전문가입니다." if lang == "한국어" else "You are a statistical analysis auditor."},
            {"role": "user", "content": prompt}
        ]
        result = await self.make_openai_call(messages)
        result_str = result.strip() if hasattr(result, 'strip') else str(result)
        if result_str.lower().startswith("reject"):
            decision = "reject"
            feedback = result_str[len("reject"):].strip(": ").strip()
            hallucination_reject_num += 1
            if hasattr(state, 'revised_analysis_history'):
                state.revised_analysis_history.append(table_analysis)
            else:
                state.revised_analysis_history = [table_analysis]
        else:
            decision = "accept"
            feedback = ""
        state.hallucination_check = decision
        state.feedback = feedback
        state.hallucination_reject_num = hallucination_reject_num
        return state
    
    async def revise_table_analysis(self, state: AgentState, on_step=None) -> AgentState:
        """분석 수정 노드 (Streamlit 프롬프트/로직 반영)"""
        if on_step:
            on_step("✏️ 분석 수정 노드 시작")
        lang = getattr(state, 'lang', '한국어')
        # report_to_modify는 revised_history가 있으면 마지막 것을, 없으면 초안을 fallback
        report_to_modify = state.revised_analysis_history[-1] if getattr(state, 'revised_analysis_history', None) else getattr(state, 'table_analysis', '')
        prompt = self.REVISION_PROMPT[lang].format(
            linearized_table=getattr(state, 'linearized_table', ''),
            ft_test_summary=str(getattr(state, 'ft_test_summary', '')),
            anchor=getattr(state, 'anchor', ''),
            report_to_modify=report_to_modify,
            feedback=getattr(state, 'feedback', '')
        )
        messages = [
            {"role": "system", "content": "당신은 통계 데이터를 바탕으로 인구집단 간 패턴과 경향성을 객관적으로 요약하는 데이터 분석 전문가입니다." if lang == "한국어" else "You are a data analyst who objectively summarizes population-level patterns based on statistical data."},
            {"role": "user", "content": prompt}
        ]
        result = await self.make_openai_call(messages)
        new_revised_analysis = result.strip() if hasattr(result, 'strip') else str(result)
        # Append to revision history
        if hasattr(state, 'revised_analysis_history') and state.revised_analysis_history:
            revision_history = state.revised_analysis_history
        else:
            revision_history = []
        revision_history.append(new_revised_analysis)
        state.revised_analysis = new_revised_analysis
        state.revised_analysis_history = revision_history
        return state
    
    async def sentence_polish_node(self, state: AgentState, on_step=None) -> AgentState:
        """문장 다듬기 노드 (Streamlit 프롬프트/로직 반영)"""
        if on_step:
            on_step("💅 문장 다듬기 노드 시작")
        lang = getattr(state, 'lang', '한국어')
        hallucination_reject_num = getattr(state, 'hallucination_reject_num', 0)
        raw_summary = state.revised_analysis if hallucination_reject_num > 0 else state.table_analysis
        prompt = self.POLISHING_PROMPT[lang].format(raw_summary=raw_summary)
        messages = [
            {"role": "system", "content": "당신은 통계 데이터를 바탕으로 작성된 한국어 보고서를 다듬는 문체 전문 에디터입니다." if lang == "한국어" else "You are a stylistic editor for statistical summaries written in Korean."},
            {"role": "user", "content": prompt}
        ]
        result = await self.make_openai_call(messages)
        polishing_result = result.strip() if hasattr(result, 'strip') else str(result)
        state.polishing_result = polishing_result
        return state
    
    async def execute(self, file_content: bytes, file_name: str, options: Dict[str, Any] = None, raw_data_content: bytes = None, raw_data_filename: str = None) -> Dict[str, Any]:
        """워크플로우 실행"""
        try:
            # 초기 상태 설정
            state = AgentState(
                uploaded_file=file_content,
                file_path=file_name,
                analysis_type=options.get("analysis_type", True) if options else True,
                selected_key=options.get("selected_key", "") if options else "",
                lang=options.get("lang", "한국어") if options else "한국어",
                user_id=options.get("user_id") if options else None,
                raw_data_file=raw_data_content
            )
            
            on_step = options.get("on_step") if options else None
            
            # 워크플로우 실행
            state = await self.table_parser_node(state, on_step)
            state = await self.hypothesis_generate_node(state, on_step)
            state = await self.test_decision_node(state, on_step)
            state = await self.ft_analysis_node(state, on_step)
            state = await self.get_anchor_node(state, on_step)
            state = await self.table_analyzer(state, on_step)
            
            # 환각 검증 및 수정 루프
            max_revisions = 4
            while state.hallucination_reject_num < max_revisions:
                state = await self.hallucination_check_node(state, on_step)
                if state.hallucination_check == "accept":
                    break
                elif state.hallucination_check == "reject":
                    if state.hallucination_reject_num >= 4:
                        if on_step:
                            on_step("⚠️ 거부 횟수 초과, 종료합니다.")
                        break
                    state = await self.revise_table_analysis(state, on_step)
                    state.hallucination_reject_num += 1
                else:
                    raise Exception(f"예상치 못한 결정: {state.hallucination_check}")
            
            state = await self.sentence_polish_node(state, on_step)
            
            return {
                "success": True,
                "result": {
                    "polishing_result": state.polishing_result,
                    "table_analysis": state.table_analysis,
                    "ft_test_summary": state.ft_test_summary,
                    "generated_hypotheses": state.generated_hypotheses,
                    "anchor": state.anchor,
                    "revised_analysis_history": state.revised_analysis_history,
                    "test_type": state.test_type,
                    "ft_test_result": state.ft_test_result.to_dict(orient="records") if hasattr(state.ft_test_result, "to_dict") else state.ft_test_result,
                    "revised_analysis_history": state.revised_analysis_history
                }
            }
            
        except Exception as e:
            return {
                "success": False,
                "error": str(e)
            } 

    async def execute_batch(self, file_content: bytes, file_name: str, test_type_map: dict, lang: str = "한국어", user_id: Optional[str] = None, raw_data_content: Optional[bytes] = None, raw_data_filename: Optional[str] = None) -> dict:
        """전체(배치) 분석: 질문별 test_type을 받아 각 질문에 대해 분석을 수행"""
        from copy import deepcopy
        # 테이블 파싱
        parsed = await self.load_survey_tables(file_content, file_name)
        tables = parsed["tables"]
        question_texts = parsed["question_texts"]
        question_keys = parsed["question_keys"]
        results = {}
        for key in question_keys:
            state = AgentState(
                uploaded_file=file_content,
                file_path=file_name,
                analysis_type=False,
                selected_key=key,
                lang=lang,
                user_id=user_id,
                raw_data_file=raw_data_content
            )
            # 테이블/질문 세팅
            state.tables = tables
            state.question_texts = question_texts
            state.question_keys = question_keys
            state.selected_key = key
            state.selected_table = tables[key]
            state.selected_question = question_texts[key]
            # test_type 지정 (LLM 추천 or 사용자 수정)
            state.test_type = test_type_map.get(key, "ft_test")
            # 워크플로우 실행 (가설~폴리싱까지)
            try:
                state = await self.hypothesis_generate_node(state)
                state = await self.ft_analysis_node(state)
                state = await self.get_anchor_node(state)
                state = await self.table_analyzer(state)
                # 환각 검증 및 수정 루프
                max_revisions = 4
                while state.hallucination_reject_num < max_revisions:
                    state = await self.hallucination_check_node(state)
                    if state.hallucination_check == "accept":
                        break
                    elif state.hallucination_check == "reject":
                        if state.hallucination_reject_num >= 4:
                            break
                        state = await self.revise_table_analysis(state)
                        state.hallucination_reject_num += 1
                    else:
                        break
                state = await self.sentence_polish_node(state)
                results[key] = {
                    "polishing_result": state.polishing_result,
                    "table_analysis": state.table_analysis,
                    "ft_test_summary": state.ft_test_summary,
                    "generated_hypotheses": state.generated_hypotheses,
                    "anchor": state.anchor,
                    "revised_analysis_history": state.revised_analysis_history,
                    "test_type": state.test_type,
                    "ft_test_result": state.ft_test_result.to_dict(orient="records") if hasattr(state.ft_test_result, "to_dict") else state.ft_test_result,
                }
            except Exception as e:
                results[key] = {"error": str(e)}
        return {"success": True, "result": results} 